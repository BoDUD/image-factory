"""Untrusted customer input is always data, never executable code."""
import csv
import hashlib
import json
import re
from pathlib import Path

MAX_ROWS = 10000
MAX_BYTES = 20 * 1024 * 1024


def digest(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def stable_hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=True, sort_keys=True).encode()).hexdigest()


def sheets(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def read_customers(path, sheet=None, encoding="utf-8-sig"):
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("名单文件超过 20 MB")
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError("请先选择 Excel 工作表")
            ws = book[sheet] if sheet else book.active
            values = []
            for i, row in enumerate(ws.iter_rows()):
                if i > MAX_ROWS:
                    raise ValueError("名单最多 10000 行")
                if any(c.data_type == "f" for c in row):
                    raise ValueError(f"第 {i + 1} 行含公式，请在 Excel 中粘贴为值再导入")
                if any(c.data_type == "e" for c in row):
                    raise ValueError(f"第 {i + 1} 行含 Excel 错误值")
                values.append(["" if c.value is None else str(c.value) for c in row])
        finally:
            book.close()
    elif path.suffix.lower() in {".csv", ".txt"}:
        try:
            text = path.read_text(encoding=encoding)
        except UnicodeDecodeError as e:
            raise ValueError("编码不匹配，请选择 GB18030 或另存为 UTF-8") from e
        values = [["customer"]] + [[s] for s in text.splitlines() if s.strip()] if path.suffix.lower() == ".txt" else list(csv.reader(text.splitlines(keepends=True)))
    else:
        raise ValueError("支持 XLSX、CSV、TXT 名单")
    if len(values) < 2:
        raise ValueError("名单需要表头及至少一条客户信息")
    headers = [s.strip() for s in values[0]]
    while headers and not headers[-1]:
        headers.pop()
    if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("表头不能为空或重复")
    if "__index" in headers:
        raise ValueError("__index 为内部保留字段")
    rows = []
    for line, row in enumerate(values[1:], 2):
        if not any(str(v).strip() for v in row):
            continue
        if len(row) > len(headers) and any(row[len(headers):]):
            raise ValueError(f"第 {line} 行的列数超过表头")
        row = row[:len(headers)] + [""] * max(0, len(headers) - len(row))
        rows.append(dict(zip(headers, row), __index=f"{len(rows) + 1:04d}"))
        if len(rows) > MAX_ROWS:
            raise ValueError("名单最多 10000 行")
    if not rows:
        raise ValueError("没有有效客户记录")
    return headers, rows


def safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip().rstrip(". ")
    value = value[:80] or "unnamed"
    if re.fullmatch(r"(?i)(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?", value):
        value = "_" + value
    return value


def within(root, relative):
    root = Path(root).resolve()
    target = (root / relative).resolve()
    if not target.is_relative_to(root) or target == root:
        raise ValueError("输出路径超出指定目录")
    return target


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def auto_bind(fields, headers):
    aliases = {"customer": ["姓名", "客户", "客户姓名", "署名", "name"], "order_id": ["编号", "订单号", "订单编号"], "text": ["文案", "文字", "内容"]}
    result = {}
    for field in fields:
        candidates = [h for h in headers if h.casefold() == field.casefold() or h in aliases.get(field, [])]
        if len(candidates) == 1:
            result[field] = candidates[0]
    return result
